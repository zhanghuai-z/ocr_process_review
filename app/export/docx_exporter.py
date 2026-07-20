"""DOCX 导出：python-docx，按块段落输出。"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.export.base import ExporterBase
from app.export.rendering import RichReflowBlock, rich_reflow_pages
from app.export.ir_builder import build_export_ir
from app.models.export_snapshot import ExportProjectSnapshot


class DocxExporter(ExporterBase):

    def export(self, snapshot: ExportProjectSnapshot, out_path: str) -> None:
        document = build_export_ir(snapshot, "docx")
        doc = Document()
        doc.core_properties.title = snapshot.project.name
        doc.core_properties.author = "OCR Process"

        for page, blocks in rich_reflow_pages(document):
            doc.add_heading(f"第 {page.page_number} 页", level=1)

            for reflow in blocks:
                self._add_block(doc, reflow)

            doc.add_page_break()

        doc.save(out_path)

    def _add_block(self, doc: Document, reflow: RichReflowBlock) -> None:
        if reflow.role == "heading":
            for text in reflow.lines:
                doc.add_heading(text, level=2)
            return

        for text in reflow.lines:
            para = _safe_add_paragraph(doc, _docx_style(reflow))
            if reflow.role == "equation":
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(text)
            run.font.size = Pt(_font_size(reflow))
            run.italic = reflow.role == "caption"
            if reflow.proof_status == "auto_flagged":
                run.font.color.rgb = RGBColor(0xF4, 0x43, 0x36)
            elif reflow.proof_status == "modified":
                run.font.color.rgb = RGBColor(0xFF, 0xA7, 0x26)


def _safe_add_paragraph(doc: Document, style_name: str):
    try:
        return doc.add_paragraph(style=style_name)
    except KeyError:
        return doc.add_paragraph()


def _docx_style(block: RichReflowBlock) -> str:
    if block.role == "caption":
        return "Caption"
    return "Normal"


def _font_size(block: RichReflowBlock) -> int:
    if block.role == "caption":
        return 10
    if block.role == "reference":
        return 11
    return 12
